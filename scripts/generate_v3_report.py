"""v3 筛选报告生成 — Word 文档，包含 Top 40 详情 + Bottom 10 + 统计摘要。

Usage:
  python scripts/generate_v3_report.py --screening data/processed/v3_screening
"""
from __future__ import annotations

import argparse
import csv
import os
import sys
from collections import Counter
from datetime import datetime

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from src.utils.logger import setup_logger

logger = setup_logger("generate_v3_report")


def build_report(top_csv: str, bottom_csv: str, output_path: str):
    doc = Document()

    doc.add_heading("v3 GNN 成膜预测 — 筛选报告", level=0)
    doc.add_paragraph(f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    doc.add_paragraph(
        "模型: GIN+GINE x5 + Cross-Graph Attention + 多头注意力池化\n"
        "数据: 1,635 篇文献, 2,268 训练样本 (正 544 / 负 1,724)\n"
        "筛选策略: 拓扑硬约束 + GNN 推理 + 多样性贪心选择 (同单体 <=3 次)"
    )

    with open(top_csv, "r", encoding="utf-8") as f:
        top_rows = list(csv.DictReader(f))
    with open(bottom_csv, "r", encoding="utf-8") as f:
        bottom_rows = list(csv.DictReader(f))

    # ── 统计摘要 ──
    doc.add_heading("统计摘要", level=1)
    n_train = sum(1 for r in top_rows if r["in_training_set"] == "True")
    n_commercial = sum(1 for r in top_rows
                       if r["aldehyde_source"] == "commercial" or r["amine_source"] == "commercial")
    n_new = len(top_rows) - n_train

    summary = doc.add_paragraph()
    summary.add_run(f"Top {len(top_rows)} 配对:\n")
    summary.add_run(f"  - 训练集已有: {n_train}\n")
    summary.add_run(f"  - 含商业单体: {n_commercial}\n")
    summary.add_run(f"  - 全新配对: {n_new}\n")
    summary.add_run(f"  - 概率范围: {float(top_rows[0]['gnn_prob']):.4f} ~ "
                    f"{float(top_rows[-1]['gnn_prob']):.4f}")

    topo_dist = Counter(r["topology"] for r in top_rows)
    summary.add_run(f"\n  拓扑分布: {dict(topo_dist)}")

    # ── Top 40 ──
    doc.add_heading(f"Top {len(top_rows)} 推荐配对", level=1)

    table = doc.add_table(rows=1, cols=8)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["排名", "醛单体", "胺单体", "概率", "拓扑", "来源", "训练集", "官能团"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
        for p in table.rows[0].cells[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(8)

    for idx, r in enumerate(top_rows):
        row = table.add_row()
        prob = float(r["gnn_prob"])
        ald_name = r["aldehyde_name"][:40]
        amine_name = r["amine_name"][:40]
        source = f"{r['aldehyde_source'][:1]}+{r['amine_source'][:1]}"
        in_train = "是" if r["in_training_set"] == "True" else "新"
        fg = f"C{r['aldehyde_fg']}+C{r['amine_fg']}"
        cells = [str(idx + 1), ald_name, amine_name, f"{prob:.4f}",
                 r["topology"], source, in_train, fg]
        for i, c in enumerate(cells):
            row.cells[i].text = c
            for p in row.cells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(7)
                    if prob >= 0.9:
                        run.font.color.rgb = RGBColor(0, 128, 0)

    # ── Bottom 10 ──
    doc.add_heading("Bottom 10 — 最不可能成膜配对", level=1)
    doc.add_paragraph("用于诊断模型负向信号，人工校验是否合理。")

    btable = doc.add_table(rows=1, cols=6)
    btable.style = "Light Grid Accent 1"
    bheaders = ["排名", "醛单体", "胺单体", "概率", "拓扑", "官能团"]
    for i, h in enumerate(bheaders):
        btable.rows[0].cells[i].text = h
        for p in btable.rows[0].cells[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(8)

    for idx, r in enumerate(bottom_rows):
        row = btable.add_row()
        prob = float(r["gnn_prob"])
        ald_name = r["aldehyde_name"][:40]
        amine_name = r["amine_name"][:40]
        fg = f"C{r['aldehyde_fg']}+C{r['amine_fg']}"
        cells = [str(idx + 1), ald_name, amine_name, f"{prob:.4f}", r["topology"], fg]
        for i, c in enumerate(cells):
            row.cells[i].text = c
            for p in row.cells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(7)

    # ── 模型信息 ──
    doc.add_heading("模型架构", level=1)
    doc.add_paragraph(
        "编码器: GIN+GINE x5 + JK-Net mean + 残差 + LayerNorm (1.34M)\n"
        "交互层: 双向多头 Cross-Graph Attention (264K)\n"
        "池化层: 4-query 多头注意力池化 + e_pair 原子对交互 (922K)\n"
        "预测头: FilmHead 两层 MLP (1024->256->128->1) + ConditionHead 5 任务 (669K)\n"
        "总参数: 3.19M\n"
        "损失函数: Focal Loss (alpha=0.75, gamma=2) + 条件 CE (lambda=0.1) + 化学正则化 (3条, lambda=0.005)\n"
        "训练策略: 文献级 GroupKFold (5折x3重复) + AdamW + cosine annealing + 早停"
    )

    doc.add_heading("数据管线", level=1)
    doc.add_paragraph(
        "数据库1 (xincailiao): 894篇 -> 2,644实验组\n"
        "数据库2 (shujuku2): 825篇 -> ~1,867实验组\n"
        "合并去重: 1,635篇, 4,435实验组\n"
        "过滤链: imine-only -> SMILES完整 -> 排金属 -> is_film非空\n"
        "最终训练集: 2,268组 (正544/负1,724, 正样本率24.0%)\n"
        "SMILES修复: sanitize=False降级(+220组) + 名称匹配回填(+19组)"
    )

    doc.add_heading("化学正则化", level=1)
    doc.add_paragraph(
        "3条化学硬事实规则 (仅保留几乎无例外的):\n"
        "  1. 苯环必需: (1 - min(n_phenyl, 1)) x prob\n"
        "  2. 官能团对称: max(0, 0.5 - fg_symmetry) x prob\n"
        "  3. C2对位: is_c2 x (1 - is_para) x prob\n"
        "设计哲学: 信任模型。其余规则(取代基/溶解性/刚性/共轭/位阻)交给GNN自己学。"
    )

    doc.save(output_path)
    logger.info(f"报告已保存: {output_path}")


def main():
    parser = argparse.ArgumentParser(description="v3 筛选报告生成")
    parser.add_argument("--screening", type=str, default="data/processed/v3_screening")
    parser.add_argument("--output", type=str, default=None)
    args = parser.parse_args()

    top_csv = os.path.join(args.screening, "Top40_v3.csv")
    bottom_csv = os.path.join(args.screening, "Bottom10_v3.csv")
    output = args.output or os.path.join(args.screening, "v3_Screening_Report.docx")

    if not os.path.exists(top_csv):
        logger.error(f"Top CSV 不存在: {top_csv}")
        sys.exit(1)

    build_report(top_csv, bottom_csv, output)


if __name__ == "__main__":
    main()
