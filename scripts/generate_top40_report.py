"""生成 Route A Top 40 可视化 Word 报告 — 修订版。

三模块结构:
  模块一: 数据概览
  模块二: Top 40 单体对详细信息 (结构式/SMILES/英文学名/来源/筛选依据)
  模块三: 统计分析

格式: 1.5 倍行距, 页边距 2.2cm
"""
import argparse
import io
import json
import os
import sys
from collections import Counter, defaultdict

import numpy as np
import pandas as pd
from rdkit import Chem
from rdkit.Chem import Draw, Descriptors

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from src.utils.logger import setup_logger

logger = setup_logger("top40_report")

# ── 常量 ──
CM_TO_INCHES = 0.3937
MARGIN_CM = 2.2
MARGIN_INCHES = MARGIN_CM * CM_TO_INCHES  # ≈ 0.866


# ═══════════════════════════════════════════════════
# 工具函数
# ═══════════════════════════════════════════════════

def _mol_image(smi: str, label: str = "", size=(350, 200)):
    """单分子结构图。"""
    mol = Chem.MolFromSmiles(smi)
    if mol is None:
        return None
    return Draw.MolToImage(mol, legend=label[:60], size=size)


def _add_image(doc, img, width=3.0):
    from docx.shared import Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    if img is None:
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_para_spacing(p)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    p.add_run().add_picture(buf, width=Inches(width))


def _set_para_spacing(para, line_spacing=1.5):
    """设置段落为 1.5 倍行距。"""
    from docx.shared import Pt
    para.paragraph_format.line_spacing = line_spacing


def _add_heading(doc, text, level=1):
    """添加标题并设置行距。"""
    h = doc.add_heading(text, level=level)
    _set_para_spacing(h, 1.5)
    return h


def _add_para(doc, text="", bold=False, font_size=10):
    """添加段落并设置行距。"""
    from docx.shared import Pt
    p = doc.add_paragraph()
    _set_para_spacing(p, 1.5)
    if text:
        run = p.add_run(text)
        run.font.size = Pt(font_size)
        run.bold = bold
    return p


def _add_table(doc, headers, rows, col_widths=None, font_size=8):
    from docx.shared import Inches, Pt

    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.autofit = True

    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = str(h)
        for p in cell.paragraphs:
            p.alignment = 1
            _set_para_spacing(p, 1.5)
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)

    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.rows[i + 1].cells[j]
            cell.text = str(val) if val is not None else "-"
            for p in cell.paragraphs:
                p.alignment = 1
                _set_para_spacing(p, 1.5)
                for run in p.runs:
                    run.font.size = Pt(font_size)

    if col_widths:
        for i, w in enumerate(col_widths):
            for row_obj in table.rows:
                row_obj.cells[i].width = Inches(w)
    return table


def _get_english_name(smi: str, name_lookup: dict) -> str:
    """从预生成名称字典获取英文学名。"""
    if smi in name_lookup:
        return str(name_lookup[smi])
    mol = Chem.MolFromSmiles(smi)
    if mol:
        from rdkit.Chem import rdMolDescriptors
        return rdMolDescriptors.CalcMolFormula(mol)
    return smi


def _get_lit_context(smi: str, lit_map: dict) -> list:
    """获取单体在文献中的出现记录。"""
    return lit_map.get(smi, [])


def _get_pool_info(smi: str, pool_lookup: dict) -> dict:
    """获取单体在池中的完整信息。"""
    return pool_lookup.get(smi, {})


def _build_rationale(row) -> str:
    """为每对单体生成筛选依据。"""
    reasons = []

    score = row["adjusted_score"]
    gnn = row["gnn_norm"]
    xgb = row["xgb_norm"]
    div = row["divergence"]
    topo = row.get("topology", "")
    ptype = row.get("pair_type", "")
    ald_topo = row.get("aldehyde_topo", "")
    am_topo = row.get("amine_topo", "")
    ald_hetero = row.get("ald_has_heterocycle", False)
    am_hetero = row.get("am_has_heterocycle", False)

    # 1. 模型共识度
    if div < 0.05:
        reasons.append(f"双模型高度一致 (分歧={div:.2f})")
    elif div < 0.15:
        reasons.append(f"双模型基本一致 (分歧={div:.2f})")
    elif div < 0.30:
        reasons.append(f"双模型存在一定分歧 (分歧={div:.2f})，GNN 占主导")
    else:
        reasons.append(f"双模型分歧较大 (分歧={div:.2f})，综合分由分歧惩罚调节")

    # 2. GNN vs XGB 贡献
    if gnn > 0.85 and xgb > 0.70:
        reasons.append("GNN 和 XGBoost 均给出高分，预测置信度高")
    elif gnn > 0.85:
        reasons.append(f"GNN 高分 ({gnn:.3f}) 驱动，XGB 辅助 ({xgb:.3f})")
    elif xgb > 0.80:
        reasons.append(f"XGBoost 高分 ({xgb:.3f}) 驱动，GNN 辅助 ({gnn:.3f})")

    # 3. 拓扑匹配
    if "六方" in str(topo):
        if am_topo == "C3" and ald_topo == "C2":
            reasons.append("经典「大胺小醛」六方拓扑 — C3 胺提供三角形节点，C2 醛为直线连接臂")
        elif ald_topo == "C3" and am_topo == "C2":
            reasons.append("「大醛小胺」六方拓扑 — C3 醛提供三角形节点，C2 胺为直线连接臂")
        elif ald_topo == "C3" and am_topo == "C3":
            reasons.append("双 C3 六方拓扑 — 两个三角单体共聚，网孔更小")
    elif "四方" in str(topo):
        reasons.append("四方 (sql) 拓扑 — 双 C2/C4 单体形成方格网格")

    # 4. 氟策略
    if "F-" in str(ptype):
        reasons.append("含氟配对 — 氟原子可改善薄膜疏水性和结晶度")

    # 5. 杂环
    if not ald_hetero and not am_hetero:
        reasons.append("双单体纯苯环骨架，无杂环干扰")

    # 6. 综合
    if score >= 0.90:
        reasons.append(f"综合得分 {score:.3f} 位于第一梯队")
    elif score >= 0.85:
        reasons.append(f"综合得分 {score:.3f} 排名靠前")

    return "；".join(reasons) + "。"


# ═══════════════════════════════════════════════════
# 主函数
# ═══════════════════════════════════════════════════

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", default="data/processed/route_a_gnn_top40_with_hard.csv")
    parser.add_argument("--full", default="data/processed/route_a_gnn_top40_with_hard_full.csv")
    parser.add_argument("--pool", default="data/processed/merged_monomer_pool.csv")
    parser.add_argument("--lit-map", default="data/processed/smiles_lit_context.json")
    parser.add_argument("--name-file", default="data/processed/monomer_names.json")
    parser.add_argument("--output", default="data/processed/Top40_Report_v2.docx")
    args = parser.parse_args()

    from docx import Document
    from docx.shared import Inches, Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    # ── 加载数据 ──
    top = pd.read_csv(args.input, encoding="utf-8-sig")
    full = pd.read_csv(args.full, encoding="utf-8-sig") if os.path.exists(args.full) else None

    # 单体名称字典: canonical_smiles → English name
    with open(args.name_file, encoding="utf-8") as f:
        name_lookup = json.load(f)

    # 单体池索引: canonical_smiles → row (for source info)
    pool = pd.read_csv(args.pool)
    pool_lookup = {}
    for _, r in pool.iterrows():
        smi = str(r["smiles"])
        mol = Chem.MolFromSmiles(smi)
        if mol:
            can = Chem.MolToSmiles(mol, canonical=True)
            pool_lookup[can] = {k: str(r[k]) if pd.notna(r[k]) else "" for k in r.index}

    # 文献上下文索引
    with open(args.lit_map, encoding="utf-8") as f:
        lit_map = json.load(f)

    # ── 创建文档 ──
    doc = Document()

    # 页边距: 上下左右 2.2 cm
    for section in doc.sections:
        section.top_margin = Cm(MARGIN_CM)
        section.bottom_margin = Cm(MARGIN_CM)
        section.left_margin = Cm(MARGIN_CM)
        section.right_margin = Cm(MARGIN_CM)

    # 默认段落样式: 1.5 倍行距
    style = doc.styles["Normal"]
    style.paragraph_format.line_spacing = 1.5

    # ═══════════════════════════════════════════════
    # 封面
    # ═══════════════════════════════════════════════
    title = doc.add_heading("Route A Top 40 筛选报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_para_spacing(title, 1.5)
    p = _add_para(doc, "", font_size=12)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("GNN + XGBoost 集成 · 化学硬规则 · C3 双模分层").font.size = Pt(12)

    n_monomers = len(pool) if full is None else len(
        set(list(full.get("aldehyde_smiles", [])) + list(full.get("amine_smiles", []))))
    n_pairs = len(full) if full is not None else 0
    p = _add_para(doc, "", font_size=10)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"单体池: {n_monomers} 可用 | 候选对: {n_pairs:,} | 模型: GNN (v4) + XGBoost").font.size = Pt(10)
    doc.add_page_break()

    # ═══════════════════════════════════════════════
    # 模块一: 数据概览
    # ═══════════════════════════════════════════════
    _add_heading(doc, "模块一：数据概览", level=1)

    if full is not None:
        n_ald = full["aldehyde_topo"].nunique() if "aldehyde_topo" in full.columns else "?"
        n_am = full["amine_topo"].nunique() if "amine_topo" in full.columns else "?"
        # Count unique SMILES
        ald_unique = full["aldehyde_smiles"].nunique()
        am_unique = full["amine_smiles"].nunique()
        n_hex = (full["topology"].str.startswith("六方")).sum() if "topology" in full.columns else 0
        n_sql = (full["topology"].str.startswith("四方")).sum() if "topology" in full.columns else 0
        n_nonstd = len(full) - n_hex - n_sql
        n_high_div = (full["divergence"] > 0.5).sum() if "divergence" in full.columns else 0
        n_hetero = (full.get("ald_has_heterocycle", pd.Series([False] * len(full))) |
                    full.get("am_has_heterocycle", pd.Series([False] * len(full)))).sum()
        n_chain_penalized = (full.get("chain_penalty", pd.Series([1.0] * len(full))) < 1.0).sum()
    else:
        ald_unique, am_unique, n_hex, n_sql, n_nonstd, n_high_div, n_hetero, n_chain_penalized = "?", "?", "?", "?", "?", "?", "?", "?"

    overview = [
        ["可用单体", f"{n_monomers} (醛={ald_unique}, 胺={am_unique})"],
        ["全量配对", f"{n_pairs:,}"],
        ["六方 (hcb) 拓扑", f"{n_hex:,} ({n_hex/len(full)*100:.0f}%)" if full is not None else "-"],
        ["四方 (sql) 拓扑", f"{n_sql:,} ({n_sql/len(full)*100:.0f}%)" if full is not None else "-"],
        ["非标准拓扑", f"{n_nonstd:,}" if full is not None else "-"],
        ["含杂环配对", f"{n_hetero:,}" if full is not None else "-"],
        ["高分歧 (>0.5)", f"{n_high_div} / {len(full):,}" if full is not None else "-"],
        ["当前规则配置", "规则#1芳环≤4 [关闭] | 规则#3杂环降权 [关闭]"],
        ["对称性判断", "CanonicalRankAtoms 全分子拓扑对称 (中心+镜像), 无 Morgan 回退"],
        ["C2 对位", "单苯环→1,4对位 | 多苯环→对称+非对位 | ≥3直链苯→软惩罚"],
        ["直链苯惩罚", f"已影响 {int(n_chain_penalized):,} 对 (≥3对位直链苯→软惩罚)" if full is not None else "已启用"],
    ]
    _add_table(doc, ["指标", "值"], overview, col_widths=[2.5, 4.5], font_size=9)
    doc.add_page_break()

    # ═══════════════════════════════════════════════
    # 模块二: Top 40 单体对详细信息
    # ═══════════════════════════════════════════════
    _add_heading(doc, "模块二：Top 40 单体对详细信息", level=1)

    for i, (_, row) in enumerate(top.iterrows()):
        rank = i + 1
        ald_smi = str(row["aldehyde_smiles"])
        am_smi = str(row["amine_smiles"])
        score = row["adjusted_score"]
        gnn = row["gnn_norm"]
        xgb = row["xgb_norm"]
        div = row["divergence"]
        topo = row.get("topology", "")
        ptype = row.get("pair_type", "")

        # 标题
        _add_heading(doc, f"#{rank} — 综合分 {score:.3f} | {topo} | {ptype}", level=2)

        # ─ 分数表格 ─
        _add_table(doc, ["指标", "值"],
                   [["综合分", f"{score:.3f}"],
                    ["GNN 归一化分", f"{gnn:.3f}"],
                    ["XGBoost 归一化分", f"{xgb:.3f}"],
                    ["分歧度", f"{div:.2f}"],
                    ["拓扑类型", str(topo)],
                    ["氟策略", str(ptype)]],
                   col_widths=[2.0, 4.0], font_size=9)

        _add_para(doc)

        # ─ 结构式图片 ─
        ald_img = _mol_image(ald_smi, f"Aldehyde: {_get_english_name(ald_smi, name_lookup)[:50]}")
        am_img = _mol_image(am_smi, f"Amine: {_get_english_name(am_smi, name_lookup)[:50]}")

        # 醛信息
        _add_heading(doc, "醛单体", level=3)
        _add_image(doc, ald_img, width=3.5)

        ald_name = _get_english_name(ald_smi, name_lookup)
        ald_pool = _get_pool_info(ald_smi, pool_lookup)
        ald_lit = _get_lit_context(ald_smi, lit_map)

        _add_table(doc, ["属性", "值"],
                   [["SMILES", ald_smi],
                    ["英文学名", ald_name],
                    ["拓扑", str(row.get("aldehyde_topo", "?"))],
                    ["含氟", str(row.get("aldehyde_f", "?"))],
                    ["来源类型", str(ald_pool.get("source", "?"))],
                    ["文献出现次数", str(len(ald_lit))],
                    ["商业编号", str(ald_pool.get("commercial_id", "-")) if ald_pool.get("commercial_id", "") not in ("", "nan") else "-"]],
                   col_widths=[1.8, 4.5], font_size=9)

        # 醛的文献来源
        if ald_lit:
            lit_ald_rows = []
            for ctx in ald_lit[:5]:  # 最多 5 条
                lit_ald_rows.append([
                    ctx.get("lit_id", "")[:60],
                    "成膜" if ctx.get("label") == 1 else "不成膜",
                    ctx.get("has_f", "?"),
                    ctx.get("solvent", "")[:80],
                    ctx.get("temperature", "")[:40],
                ])
            _add_para(doc, "文献来源 (该单体出现记录，最多展示 5 条):", font_size=9)
            _add_table(doc, ["文献 ID", "成膜", "含氟", "溶剂", "温度"],
                       lit_ald_rows, col_widths=[2.5, 0.6, 0.5, 2.0, 1.2], font_size=7)

        _add_para(doc)

        # 胺信息
        _add_heading(doc, "胺单体", level=3)
        _add_image(doc, am_img, width=3.5)

        am_name = _get_english_name(am_smi, name_lookup)
        am_pool = _get_pool_info(am_smi, pool_lookup)
        am_lit = _get_lit_context(am_smi, lit_map)

        _add_table(doc, ["属性", "值"],
                   [["SMILES", am_smi],
                    ["英文学名", am_name],
                    ["拓扑", str(row.get("amine_topo", "?"))],
                    ["含氟", str(row.get("amine_f", "?"))],
                    ["来源类型", str(am_pool.get("source", "?"))],
                    ["文献出现次数", str(len(am_lit))],
                    ["商业编号", str(am_pool.get("commercial_id", "-")) if am_pool.get("commercial_id", "") not in ("", "nan") else "-"]],
                   col_widths=[1.8, 4.5], font_size=9)

        # 胺的文献来源
        if am_lit:
            lit_am_rows = []
            for ctx in am_lit[:5]:
                lit_am_rows.append([
                    ctx.get("lit_id", "")[:60],
                    "成膜" if ctx.get("label") == 1 else "不成膜",
                    ctx.get("has_f", "?"),
                    ctx.get("solvent", "")[:80],
                    ctx.get("temperature", "")[:40],
                ])
            _add_para(doc, "文献来源 (该单体出现记录，最多展示 5 条):", font_size=9)
            _add_table(doc, ["文献 ID", "成膜", "含氟", "溶剂", "温度"],
                       lit_am_rows, col_widths=[2.5, 0.6, 0.5, 2.0, 1.2], font_size=7)

        _add_para(doc)

        # ─ 筛选依据 ─
        _add_heading(doc, "筛选依据", level=3)
        rationale = _build_rationale(row)
        _add_para(doc, rationale, font_size=10)

        # 页内分页 (最后一对不加)
        if rank < len(top):
            doc.add_page_break()

    doc.add_page_break()

    # ═══════════════════════════════════════════════
    # 模块三: 统计分析
    # ═══════════════════════════════════════════════
    _add_heading(doc, "模块三：统计分析", level=1)

    # 3.1 拓扑分布
    _add_heading(doc, "3.1 拓扑分布 (Top 40)", level=2)
    topo_counts = top["topology"].value_counts()
    _add_table(doc, ["拓扑", "数量", "占比"],
               [[t, c, f"{c/40*100:.0f}%"] for t, c in topo_counts.items()],
               col_widths=[2.0, 1.0, 1.0], font_size=9)

    # 3.2 氟策略分布
    _add_heading(doc, "3.2 氟策略分布", level=2)
    f_counts = top["pair_type"].value_counts()
    _add_table(doc, ["氟策略", "数量", "占比"],
               [[t, c, f"{c/40*100:.0f}%"] for t, c in f_counts.items()],
               col_widths=[2.5, 1.0, 1.0], font_size=9)

    # 3.3 C3 分层占比
    _add_heading(doc, "3.3 C3 分层占比", level=2)
    c3_am = (top["amine_topo"] == "C3").sum()
    c3_ald = (top["aldehyde_topo"] == "C3").sum()
    _add_table(doc, ["指标", "实际", "目标"],
               [["C3-胺 (大胺小醛)", f"{c3_am}/40 ({c3_am*100/40:.0f}%)", "45%"],
                ["C3-醛 (大醛小胺)", f"{c3_ald}/40 ({c3_ald*100/40:.0f}%)", "45%"],
                ["其余组合", f"{40-c3_am-c3_ald}/40 ({(40-c3_am-c3_ald)*100/40:.0f}%)", "10%"]],
               col_widths=[2.5, 2.0, 1.0], font_size=9)

    # 3.4 高频单体
    _add_heading(doc, "3.4 高频出现单体", level=2)
    ald_pool_names = {smi: _get_english_name(smi, name_lookup) for smi in top["aldehyde_smiles"]}
    am_pool_names = {smi: _get_english_name(smi, name_lookup) for smi in top["amine_smiles"]}

    ald_freq = Counter(top["aldehyde_smiles"])
    am_freq = Counter(top["amine_smiles"])
    freq_rows = []
    for smi, cnt in ald_freq.most_common(8):
        freq_rows.append([ald_pool_names.get(smi, smi)[:60], "醛", cnt])
    for smi, cnt in am_freq.most_common(8):
        freq_rows.append([am_pool_names.get(smi, smi)[:60], "胺", cnt])
    _add_table(doc, ["单体名称 (英文学名)", "类型", "出现次数"],
               freq_rows, col_widths=[4.0, 0.8, 1.0], font_size=8)

    # 3.5 分数分布
    _add_heading(doc, "3.5 综合分分布", level=2)
    scores = top["adjusted_score"]
    _add_table(doc, ["统计量", "值"],
               [["均值", f"{scores.mean():.3f}"],
                ["中位数", f"{scores.median():.3f}"],
                ["最小值", f"{scores.min():.3f}"],
                ["最大值", f"{scores.max():.3f}"],
                ["标准差", f"{scores.std():.3f}"]],
               col_widths=[2.0, 2.0], font_size=9)

    # 3.6 化学规则状态
    _add_heading(doc, "3.6 当前化学规则状态", level=2)
    _add_table(doc, ["规则", "状态", "说明"],
               [["#0 苯环必需", "启用", "无苯环单体排除"],
                ["#1 芳环≤4", "已关闭", "空间位阻约束 — 实验关闭"],
                ["#2 对称性", "增强", "CanonicalRankAtoms 全分子拓扑对称 (中心+镜像), 无 Morgan 回退"],
                ["#3 杂环降权", "已关闭", "实验关闭 — 杂环不加分不扣分, 保持中立"],
                ["#4 C2 对位", "增强", "单苯环→1,4对位; 多苯环→对称+允许非对位"],
                ["#4b 直链苯", "新增", "≥3 对位直链苯→软惩罚 (越长惩罚越大, 非直链不惩罚)"],
                ["#5 炔丙基醚", "启用", "醚+炔共存排除"],
                ["#6 C2 取代基", "启用", ">4 取代限卤素"],
                ["#7 官能团", "启用", "≥2 醛/胺基"]],
               col_widths=[1.3, 0.8, 4.5], font_size=9)

    # 3.7 方法说明
    _add_heading(doc, "3.7 方法说明", level=2)
    methods = [
        ["模型架构", "GNN 编码器 (v4 预训练, 256维) + BilinearHead (26维机理描述符)"],
        ["训练策略", "λ_chem=0.005 化学正则化, Focal Loss (α=0.75, γ=2.0)"],
        ["集成打分", "GNN 60% + XGBoost 40% − 0.10×分歧惩罚"],
        ["C3 分层", "C3-胺 45% + C3-醛 45% + 其余 10%"],
        ["加成系数", "C3-胺 ×1.15, C3-醛 ×1.10"],
        ["直链苯惩罚", "≥3 对位直链苯环 → max(0.4, 1−(n−3)×0.08) 惩罚因子"],
        ["去重", "InChI Key 配对去重，保留最高分"],
    ]
    _add_table(doc, ["组件", "说明"], methods, col_widths=[1.5, 5.5], font_size=9)

    # ── 保存 ──
    os.makedirs(os.path.dirname(args.output), exist_ok=True)
    doc.save(args.output)
    logger.info(f"报告已保存: {args.output}")
    print(f"报告已保存: {args.output}")


if __name__ == "__main__":
    main()
